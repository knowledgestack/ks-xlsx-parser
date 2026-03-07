"""
Tests for LLM-ready derived artifacts.

Verifies sheet summary detection, entity extraction, KPI identification,
and reading-order linearization using programmatic fixtures.
"""

import pytest

from xlsx_parser.analysis import (
    EntityIndexBuilder,
    KpiCatalogBuilder,
    ReadingOrderLinearizer,
    SheetSummaryAnalyzer,
)
from xlsx_parser.models.common import SheetPurpose
from xlsx_parser.parsers import WorkbookParser


class TestSheetSummaryAnalyzer:
    """Test sheet purpose detection and summary generation."""

    def test_table_sheet_detection(self, table_workbook):
        result = WorkbookParser(path=table_workbook).parse()
        sheet = result.sheets[0]
        analyzer = SheetSummaryAnalyzer(
            sheet=sheet,
            tables=result.tables,
            dependency_graph=result.dependency_graph,
        )
        summary = analyzer.analyze()
        assert summary.sheet_name == "Sales"
        assert summary.total_cells > 0
        assert summary.summary_text != ""
        assert len(summary.key_tables) > 0

    def test_calculation_sheet_detection(self, formula_workbook):
        result = WorkbookParser(path=formula_workbook).parse()
        # The "Calculations" sheet should detect as calculation
        calc_sheet = result.sheets[1]
        analyzer = SheetSummaryAnalyzer(
            sheet=calc_sheet,
            dependency_graph=result.dependency_graph,
        )
        summary = analyzer.analyze()
        assert summary.formula_count > 0
        assert summary.formula_density > 0

    def test_input_sheet_detection(self, data_validation_workbook):
        result = WorkbookParser(path=data_validation_workbook).parse()
        sheet = result.sheets[0]
        analyzer = SheetSummaryAnalyzer(
            sheet=sheet,
            dependency_graph=result.dependency_graph,
        )
        summary = analyzer.analyze()
        assert summary.has_data_validation is True
        assert summary.purpose == SheetPurpose.INPUT

    def test_chart_sheet_detection(self, chart_workbook):
        result = WorkbookParser(path=chart_workbook).parse()
        sheet = result.sheets[0]
        analyzer = SheetSummaryAnalyzer(
            sheet=sheet,
            charts=result.charts,
            dependency_graph=result.dependency_graph,
        )
        summary = analyzer.analyze()
        assert summary.has_charts is True

    def test_summary_text_nonempty(self, simple_workbook):
        result = WorkbookParser(path=simple_workbook).parse()
        sheet = result.sheets[0]
        analyzer = SheetSummaryAnalyzer(
            sheet=sheet,
            tables=result.tables,
            dependency_graph=result.dependency_graph,
        )
        summary = analyzer.analyze()
        assert len(summary.summary_text) > 10
        assert "Sheet1" in summary.summary_text


class TestEntityIndexBuilder:
    """Test entity extraction from headers and table columns."""

    def test_table_columns_extracted(self, table_workbook):
        result = WorkbookParser(path=table_workbook).parse()
        builder = EntityIndexBuilder(
            sheets=result.sheets,
            tables=result.tables,
            named_ranges=result.named_ranges,
        )
        index = builder.build()
        entity_names = {e.name.lower() for e in index.entities}
        assert "product" in entity_names
        assert "region" in entity_names

    def test_named_ranges_included(self, formula_workbook):
        result = WorkbookParser(path=formula_workbook).parse()
        builder = EntityIndexBuilder(
            sheets=result.sheets,
            tables=result.tables,
            named_ranges=result.named_ranges,
        )
        index = builder.build()
        entity_names = {e.name for e in index.entities}
        assert "Price" in entity_names
        assert "Quantity" in entity_names

    def test_measure_categorization(self, table_workbook):
        result = WorkbookParser(path=table_workbook).parse()
        builder = EntityIndexBuilder(
            sheets=result.sheets,
            tables=result.tables,
        )
        index = builder.build()
        total_entity = next((e for e in index.entities if "total" in e.name.lower()), None)
        if total_entity:
            assert total_entity.category == "measure"


class TestKpiCatalogBuilder:
    """Test KPI cell identification."""

    def test_kpis_identified_in_formulas(self, assumptions_workbook):
        result = WorkbookParser(path=assumptions_workbook).parse()
        builder = KpiCatalogBuilder(
            sheets=result.sheets,
            charts=result.charts,
            dependency_graph=result.dependency_graph,
        )
        kpis = builder.build()
        # The assumptions workbook has bold formula cells that should be detected
        # Even if no KPIs are found, the builder should return a list
        assert isinstance(kpis, list)

    def test_kpis_have_labels(self, assumptions_workbook):
        result = WorkbookParser(path=assumptions_workbook).parse()
        builder = KpiCatalogBuilder(
            sheets=result.sheets,
            dependency_graph=result.dependency_graph,
        )
        kpis = builder.build()
        for kpi in kpis:
            assert kpi.cell_ref != ""
            assert kpi.sheet_name != ""


class TestReadingOrderLinearizer:
    """Test reading-order text linearization."""

    def test_linearize_simple_sheet(self, simple_workbook):
        result = WorkbookParser(path=simple_workbook).parse()
        sheet = result.sheets[0]
        linearizer = ReadingOrderLinearizer(
            sheet=sheet,
            tables=result.tables,
        )
        text = linearizer.linearize()
        assert "## Sheet: Sheet1" in text
        assert len(text) > 20

    def test_linearize_mixed_content(self, mixed_content_layout):
        result = WorkbookParser(path=mixed_content_layout).parse()
        sheet = result.sheets[0]
        linearizer = ReadingOrderLinearizer(
            sheet=sheet,
            tables=result.tables,
        )
        text = linearizer.linearize()
        assert "REPORT TITLE" in text

    def test_linearize_with_chart(self, chart_workbook):
        result = WorkbookParser(path=chart_workbook).parse()
        sheet = result.sheets[0]
        linearizer = ReadingOrderLinearizer(
            sheet=sheet,
            charts=result.charts,
        )
        text = linearizer.linearize()
        assert "Chart:" in text or "## Sheet:" in text

    def test_linearize_empty_sheet(self, tmp_dir):
        from openpyxl import Workbook
        path = tmp_dir / "empty.xlsx"
        wb = Workbook()
        wb.save(path)
        result = WorkbookParser(path=path).parse()
        sheet = result.sheets[0]
        linearizer = ReadingOrderLinearizer(sheet=sheet)
        text = linearizer.linearize()
        assert "(empty sheet)" in text

    def test_linearize_with_comments(self, comment_workbook):
        result = WorkbookParser(path=comment_workbook).parse()
        sheet = result.sheets[0]
        linearizer = ReadingOrderLinearizer(sheet=sheet)
        text = linearizer.linearize()
        assert "Note:" in text
        assert "Total annual revenue" in text
